"""End-to-end smoke test for the final project backend (auth + quiz feature)."""
import io
import sys

from fastapi.testclient import TestClient

from app.main import app

PASS = []
FAIL = []

def check(name, cond, extra=""):
    (PASS if cond else FAIL).append(name)
    print(f"{'PASS' if cond else 'FAIL'}: {name} {extra}")

with TestClient(app) as c:
    # ── Auth: admin login ──
    r = c.post("/api/auth/login", json={"email": "admin@gmail.com", "password": "Raza@17582"})
    check("admin login", r.status_code == 200, str(r.status_code))
    admin_h = {"Authorization": f"Bearer {r.json()['access_token']}"}

    # ── Admin creates a teacher ──
    r = c.post("/api/users/", headers=admin_h, json={
        "first_name": "Tina", "last_name": "Teacher", "email": "tina@x.edu.pk",
        "password": "Teacher@123", "role_name": "teacher", "department": "CS",
    })
    check("admin creates teacher", r.status_code == 201, f"{r.status_code} {r.text[:200]}")
    teacher_id = r.json()["id"]

    # ── Admin creates a course for the teacher ──
    r = c.post("/api/courses/", headers=admin_h, json={
        "course_code": "CS101", "title": "Intro to CS", "description": "Basics",
        "semester": 1, "session": "26-30", "session_type": "morning", "teacher_id": teacher_id,
    })
    check("admin creates course", r.status_code == 201, f"{r.status_code} {r.text[:200]}")
    course_id = r.json()["id"]

    # ── Teacher login ──
    r = c.post("/api/auth/login", json={"email": "tina@x.edu.pk", "password": "Teacher@123"})
    check("teacher login", r.status_code == 200, str(r.status_code))
    th = {"Authorization": f"Bearer {r.json()['access_token']}"}

    # ── Teacher: parse-file preview (CSV question bank) ──
    csv_bank = (
        "Question,Option A,Option B,Option C,Option D,Answer\r\n"
        "What is 2+2?,3,4,5,6,B\r\n"
        "Capital of France?,Berlin,Paris,Rome,Madrid,B\r\n"
        "HTTP default port?,80,443,8080,3000,A\r\n"
        "SQL means?,Structured Query Language,Sequential Query Logic,Simple Query Language,Some Query Language,A\r\n"
        "1 byte = ?,6 bits,8 bits,16 bits,4 bits,B\r\n"
        "2^10 = ?,512,1024,2048,4096,B\r\n"
        "OSI layers?,5,7,4,6,B\r\n"
        "Git clone cmd?,git push,git pull,git clone,git init,C\r\n"
        "Avg quicksort?,O(n),O(n log n),O(n^2),O(log n),B\r\n"
        "RAM is?,volatile,non-volatile,permanent,write-once,A\r\n"
    ).encode()
    r = c.post("/api/quizzes/parse-file", headers=th,
               files={"file": ("bank.csv", io.BytesIO(csv_bank), "text/csv")})
    check("parse CSV bank", r.status_code == 200 and r.json()["count"] == 10, f"{r.status_code} {r.text[:200]}")

    # ── Teacher: parse-file preview (DOCX question bank) ──
    # Build a minimal docx in-memory
    import zipfile
    def make_docx(paragraphs):
        doc_xml = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
                   '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>']
        for p in paragraphs:
            doc_xml.append(f'<w:p><w:r><w:t xml:space="preserve">{p}</w:t></w:r></w:p>')
        doc_xml.append('</w:body></w:document>')
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types/>')
            zf.writestr("word/document.xml", "".join(doc_xml))
        buf.seek(0)
        return buf

    docx_buf = make_docx([
        "1. What is the capital of Pakistan?",
        "a) Karachi",
        "b) Lahore",
        "c) Islamabad",
        "d) Peshawar",
        "Answer: c",
        "2. 5 + 7 = ?",
        "a) 10",
        "b) 12",
        "c) 13",
        "d) 11",
        "Answer: b",
    ])
    r = c.post("/api/quizzes/parse-file", headers=th,
               files={"file": ("bank.docx", docx_buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    check("parse DOCX bank", r.status_code == 200 and r.json()["count"] == 2, f"{r.status_code} {r.text[:300]}")

    # ── Teacher: create quiz from CSV bank with total_questions=3 ──
    r = c.post("/api/quizzes", headers=th, data={
        "course_id": course_id, "title": "Shuffled Quiz", "description": "bank quiz",
        "total_questions": "3",
    }, files={"question_file": ("bank.csv", io.BytesIO(csv_bank), "text/csv")})
    check("create quiz from file", r.status_code == 201, f"{r.status_code} {r.text[:300]}")
    quiz_id = r.json()["id"]
    check("quiz total_questions saved", r.json()["total_questions"] == 3, str(r.json()))
    check("quiz question_count = 10", r.json()["question_count"] == 10, str(r.json()["question_count"]))

    # ── Teacher: create quiz with manual questions too ──
    r = c.post("/api/quizzes", headers=th, data={
        "course_id": course_id, "title": "Manual Quiz",
        "questions": '[{"text": "Pick one", "options": ["a1", "a2"], "correct": 1}]',
    })
    check("create manual quiz", r.status_code == 201, f"{r.status_code} {r.text[:200]}")
    manual_id = r.json()["id"]

    # ── Teacher sees full bank with correct answers ──
    r = c.get(f"/api/quizzes/{quiz_id}", headers=th)
    check("teacher gets full bank", r.status_code == 200 and len(r.json()["questions"]) == 10,
          f"{len(r.json().get('questions', []))} qs")
    check("teacher sees correct answers", all(q["correct_index"] is not None for q in r.json()["questions"]))

    # ── Register two students (register → OTP from console/DB) ──
    import re
    from app.database.database import SessionLocal
    from app.models import User, OTPVerification

    def register_student(first, last, email, roll):
        r = c.post("/api/auth/register", json={
            "first_name": first, "last_name": last, "email": email,
            "password": "Student@123", "confirm_password": "Student@123",
            "roll_number": roll, "semester": 1, "enrollment_year": 2026,
            "department": "CS", "session_type": "morning",
        })
        check(f"register {first}", r.status_code == 201, f"{r.status_code} {r.text[:200]}")
        # Read OTP straight from the DB (email provider is console in dev)
        db = SessionLocal()
        try:
            u = db.query(User).filter(User.email == email).one()
            rec = (db.query(OTPVerification)
                     .filter(OTPVerification.user_id == u.id, OTPVerification.is_used == False)  # noqa: E712
                     .order_by(OTPVerification.created_at.desc()).first())
            # otp is stored hashed — verify via service
            from app.services.otp_service import verify_otp
            # We need the plaintext; regenerate by creating a fresh one is easier:
            # instead, compute using the stored hash check helper
            from app.dependencies.auth import pwd_context
            # Brute force not possible — so create a known OTP directly
            from app.models import utcnow
            rec.otp_hash = pwd_context.hash("123456")
            rec.expires_at = utcnow().replace(hour=23, minute=59) if False else rec.expires_at
            db.commit()
        finally:
            db.close()
        r = c.post("/api/auth/verify-otp", json={"email": email, "otp_code": "123456"})
        check(f"verify OTP {first}", r.status_code == 200, f"{r.status_code} {r.text[:200]}")
        return {"Authorization": f"Bearer {r.json()['access_token']}"}

    s1 = register_student("Sami", "Student", "sami@x.edu.pk", "26-CS-01")
    s2 = register_student("Sara", "Student", "sara@x.edu.pk", "26-CS-02")

    # ── Students see the quiz in their list ──
    r = c.get("/api/quizzes", headers=s1)
    check("student lists quizzes", r.status_code == 200 and any(q["id"] == quiz_id for q in r.json()),
          f"{r.status_code} {len(r.json()) if r.status_code == 200 else r.text[:200]}")

    # ── Each student gets their own shuffled subset of 3 ──
    r1 = c.get(f"/api/quizzes/{quiz_id}", headers=s1)
    r2 = c.get(f"/api/quizzes/{quiz_id}", headers=s2)
    check("student1 gets 3 questions", r1.status_code == 200 and len(r1.json()["questions"]) == 3,
          f"{len(r1.json().get('questions', []))}")
    check("student2 gets 3 questions", r2.status_code == 200 and len(r2.json()["questions"]) == 3,
          f"{len(r2.json().get('questions', []))}")
    check("students cannot see answers",
          all(q["correct_index"] is None for q in r1.json()["questions"]))

    ids1 = [q["id"] for q in r1.json()["questions"]]
    ids2 = [q["id"] for q in r2.json()["questions"]]
    check("students got different subsets/orders", ids1 != ids2)

    # ── Refetch gives the SAME subset (deterministic) ──
    r1b = c.get(f"/api/quizzes/{quiz_id}", headers=s1)
    check("refetch is stable for same student",
          [q["id"] for q in r1b.json()["questions"]] == ids1)

    # ── Student1 submits: correct answers = all 3 (use full bank to find corrects) ──
    r = c.get(f"/api/quizzes/{quiz_id}", headers=th)
    correct_map = {q["id"]: q["correct_index"] for q in r.json()["questions"]}
    answers = [{"question_id": q["id"], "selected_index": correct_map[q["id"]]}
               for q in r1.json()["questions"]]
    r = c.post(f"/api/quizzes/{quiz_id}/submit", headers=s1, json={"answers": answers})
    check("student1 submits (all correct)", r.status_code == 200 and r.json()["score"] == 3,
          f"{r.status_code} {r.text[:300]}")

    # ── Wrong answers score correctly ──
    answers2 = [{"question_id": q["id"],
                 "selected_index": (correct_map[q["id"]] + 1) % len(q["options"])}
                for q in r2.json()["questions"]]
    r = c.post(f"/api/quizzes/{quiz_id}/submit", headers=s2, json={"answers": answers2})
    check("student2 submits (all wrong)", r.status_code == 200 and r.json()["score"] == 0,
          f"{r.status_code} {r.text[:300]}")

    # ── Double submit rejected ──
    r = c.post(f"/api/quizzes/{quiz_id}/submit", headers=s1, json={"answers": answers})
    check("duplicate attempt rejected", r.status_code == 400, str(r.status_code))

    # ── Student attempt result ──
    r = c.get(f"/api/quizzes/{quiz_id}/attempts", headers=s1)
    check("student sees own attempt", r.status_code == 200 and r.json()["score"] == 3, r.text[:200])

    # ── Teacher sees all attempts ──
    r = c.get(f"/api/quizzes/{quiz_id}/all-attempts", headers=th)
    check("teacher sees 2 attempts", r.status_code == 200 and len(r.json()) == 2, r.text[:200])

    # ── Teacher exports quiz results as Excel ──
    r = c.get(f"/api/quizzes/{quiz_id}/export", headers=th)
    rows, ok, styled = [], r.status_code == 200 and "spreadsheetml" in r.headers.get("content-type", ""), False
    if ok:
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(r.content)).active
        rows = list(ws.iter_rows(values_only=True))
        styled = bool(ws["A2"].font.bold) and str(ws["A6"].fill.fgColor.rgb).endswith("1F2A44")
    # 5 styled header rows + 1 column-header row + 2 attempt rows
    scores = sorted([rows[6][4], rows[7][4]]) if len(rows) >= 8 else []
    check("teacher exports xlsx results",
          ok and styled and len(rows) == 8 and rows[5][0] == "#" and scores == [0, 3],
          f"{r.status_code} rows={len(rows)}")
    # Student cannot export
    r = c.get(f"/api/quizzes/{quiz_id}/export", headers=s1)
    check("student cannot export", r.status_code == 403, str(r.status_code))

    # ── Validation: total_questions > bank size rejected ──
    r = c.post("/api/quizzes", headers=th, data={
        "course_id": course_id, "title": "Too many", "total_questions": "99",
    }, files={"question_file": ("bank.csv", io.BytesIO(csv_bank), "text/csv")})
    check("total_questions > bank rejected", r.status_code == 400, str(r.status_code))

    # ── Validation: empty quiz rejected ──
    r = c.post("/api/quizzes", headers=th, data={"course_id": course_id, "title": "Empty"})
    check("empty quiz rejected", r.status_code == 400, str(r.status_code))

    # ── Student cannot create quizzes ──
    r = c.post("/api/quizzes", headers=s1, data={"course_id": course_id, "title": "Nope"})
    check("student cannot create quiz", r.status_code == 403, str(r.status_code))

    # ── Teacher delete works ──
    r = c.delete(f"/api/quizzes/{manual_id}", headers=th)
    check("teacher deletes quiz", r.status_code == 200, str(r.status_code))

    # ── Admin stats endpoint (used by admin dashboard) ──
    r = c.get("/api/users/stats/dashboard", headers=admin_h)
    check("admin stats", r.status_code == 200 and r.json()["total_students"] == 2, r.text[:200])

    # ── Admin lists users + courses ──
    r = c.get("/api/users/", headers=admin_h)
    check("admin lists users", r.status_code == 200 and len(r.json()) >= 4, str(len(r.json())))
    r = c.get("/api/courses/", headers=admin_h)
    check("admin lists courses", r.status_code == 200 and len(r.json()) == 1, str(len(r.json())))

    # ── Admin hard-deletes a throwaway user (exercises cascade cleanup) ──
    register_student("Temp", "Delete", "tempdel@x.edu.pk", "26-CS-99")
    db = SessionLocal()
    try:
        temp_id = db.query(User).filter(User.email == "tempdel@x.edu.pk").one().id
    finally:
        db.close()
    r = c.delete(f"/api/users/{temp_id}/hard", headers=admin_h)
    check("admin hard-deletes user", r.status_code == 200, r.text[:200])
    r = c.get("/api/users/", headers=admin_h)
    check("hard-deleted user gone", r.status_code == 200 and temp_id not in [u["id"] for u in r.json()],
          str(len(r.json())))

    # No-trailing-slash variants (Vercel normalizes paths; both must match)
    r = c.get("/api/users", headers=admin_h)
    check("list users without trailing slash", r.status_code == 200, str(r.status_code))
    r = c.get("/api/courses", headers=admin_h)
    check("list courses without trailing slash", r.status_code == 200, str(r.status_code))

    # ── Promotion flow (admin) ──
    db = SessionLocal()
    try:
        sami_id = db.query(User).filter(User.email == "sami@x.edu.pk").one().id
        sara_id = db.query(User).filter(User.email == "sara@x.edu.pk").one().id
    finally:
        db.close()

    r = c.get("/api/users/promotion/sessions", params={"session_type": "morning"}, headers=admin_h)
    check("admin lists promotion sessions", r.status_code == 200 and any(s["session"] == "26-30" for s in r.json()),
          str(r.json())[:200])

    r = c.get("/api/users/promotion/sessions/26-30/semesters", params={"session_type": "morning"}, headers=admin_h)
    sem1 = next((s for s in r.json() if s["semester"] == 1), None) if r.status_code == 200 else None
    check("2 students in semester 1", sem1 is not None and sem1["student_count"] == 2, str(r.json())[:200])

    r = c.get("/api/users/promotion/sessions/26-30/semesters/1", params={"session_type": "morning"}, headers=admin_h)
    check("semester 1 students listed", r.status_code == 200 and len(r.json()) == 2, str(len(r.json())))

    r = c.post("/api/users/promotion/promote", json={"student_ids": [sami_id, sara_id], "to_semester": 2},
               headers=admin_h)
    check("promote 2 students to sem 2", r.status_code == 200 and r.json()["count"] == 2, r.text[:200])

    r = c.get("/api/users/promotion/history", params={"session": "26-30"}, headers=admin_h)
    check("promotion history has 2 rows", r.status_code == 200 and len(r.json()) == 2, str(r.text[:200]))

    r = c.get(f"/api/users/{sami_id}/profile", headers=admin_h)
    d = r.json() if r.status_code == 200 else {}
    check("profile: semester bumped, promotions kept, no LMS keys",
          d.get("profile", {}).get("semester") == 2
          and len(d.get("promotions", [])) == 1
          and "attendance" not in d and "assignments" not in d and "results" not in d,
          r.text[:200])

    r = c.get(f"/api/users/{sami_id}/semester-progress", headers=admin_h)
    d = r.json() if r.status_code == 200 else {}
    check("semester progress: sem1 completed, sem2 current",
          r.status_code == 200 and d["semesters"][0]["status"] == "completed"
          and d["semesters"][1]["status"] == "current",
          r.text[:200])

    r = c.post("/api/users/promotion/graduate", json={"student_ids": [sara_id]}, headers=admin_h)
    check("graduate 1 student", r.status_code == 200 and r.json()["count"] == 1, r.text[:200])
    r = c.get("/api/users/promotion/sessions/26-30/semesters/2", params={"session_type": "morning"}, headers=admin_h)
    sara_row = next((s for s in r.json() if s["id"] == sara_id), None) if r.status_code == 200 else None
    check("sara marked graduated", sara_row is not None and sara_row["is_graduated"], str(r.text[:200]))

print(f"\n{len(PASS)} passed, {len(FAIL)} failed")
if FAIL:
    print("FAILED:", FAIL)
    sys.exit(1)
