import io
import json
import random
import re
from typing import List, Optional

from datetime import datetime, timedelta, timezone

from fastapi import Request, APIRouter, Depends, HTTPException, UploadFile, File, Form, status
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.database.database import get_db
from app.dependencies.auth import get_current_user, require_teacher
from app.models import User, Quiz, QuizQuestion, QuizAttempt, QuizAttemptAnswer, QuizTimer, Course, StudentProfile
from app.schemas.quiz import (
    QuizOut, QuizDetailOut, QuestionDetailOut, ParseFileOut,
    QuizSubmitIn, QuizAttemptOut, QuizAnswerResultOut, QuizTeacherAttemptOut,
)
from app.routers.courses import student_has_access, get_student_courses
from app.dependencies.ratelimit import limiter
from app.services.question_bank import parse_question_file

router = APIRouter(prefix="/api", tags=["Quizzes"])


def _parse_options(quiz_question: QuizQuestion) -> List[str]:
    if not quiz_question.options:
        return []
    try:
        parsed = json.loads(quiz_question.options)
        return parsed if isinstance(parsed, list) else []
    except (TypeError, ValueError):
        return []


def _question_out(q: QuizQuestion, include_correct: bool, order: int = None) -> QuestionDetailOut:
    return QuestionDetailOut(
        id=q.id,
        text=q.text,
        options=_parse_options(q),
        order_index=q.order_index if order is None else order,
        correct_index=q.correct_index if include_correct else None,
    )


def _assigned_questions(quiz: Quiz, student_id: str) -> List[QuizQuestion]:
    """Deterministically pick + shuffle the questions assigned to a student.

    Seeded by quiz id + student id so every fetch and the final submission
    see the exact same set/order, while each student gets a different mix.
    """
    bank = list(quiz.questions)
    rng = random.Random(f"{quiz.id}:{student_id}")
    target = quiz.total_questions
    if target is not None and 0 < target < len(bank):
        assigned = rng.sample(bank, target)
    else:
        assigned = list(bank)
    rng.shuffle(assigned)
    return assigned


# Server-side look past the countdown absorbs network latency
TIMER_GRACE_SECONDS = 30


def _ensure_quiz_timer(db: Session, quiz: Quiz, student_id: str) -> datetime:
    """Return this student's start time, recording it on first open.

    Re-opening the quiz resumes the same countdown instead of resetting it.
    """
    timer = db.query(QuizTimer).filter(
        QuizTimer.quiz_id == quiz.id,
        QuizTimer.student_id == student_id,
    ).first()
    if timer:
        return timer.started_at
    timer = QuizTimer(quiz_id=quiz.id, student_id=student_id)
    db.add(timer)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        timer = db.query(QuizTimer).filter(
            QuizTimer.quiz_id == quiz.id,
            QuizTimer.student_id == student_id,
        ).one()
        return timer.started_at
    db.refresh(timer)
    return timer.started_at


def _quiz_out(q: Quiz) -> QuizOut:
    course = q.course
    return QuizOut(
        id=q.id,
        course_id=q.course_id,
        course_title=course.title if course else None,
        course_code=course.course_code if course else None,
        title=q.title,
        description=q.description,
        time_limit=q.time_limit,
        deadline=q.deadline,
        total_questions=q.total_questions,
        question_count=len(q.questions),
        max_marks=q.max_marks,
        is_published=q.is_published,
        created_at=q.created_at,
    )


def _score_answers(questions: List[QuizQuestion], answers, require_all: bool = True):
    """Grade the student's answers against the questions they were assigned.

    require_all=False is used for countdown auto-submits: unanswered
    questions are simply not counted as correct.
    """
    allowed = {q.id: q for q in questions}
    validated_answers = []
    seen_questions = set()
    for answer in answers:
        question = allowed.get(answer.question_id)
        if not question or answer.question_id in seen_questions:
            raise HTTPException(status_code=400, detail="Each assigned question must be answered exactly once")
        options = _parse_options(question)
        if answer.selected_index < 0 or answer.selected_index >= len(options):
            raise HTTPException(status_code=400, detail="Invalid answer selected")
        seen_questions.add(answer.question_id)
        validated_answers.append((answer, question))
    if require_all and len(seen_questions) != len(allowed):
        raise HTTPException(status_code=400, detail="Please answer all quiz questions")

    score = 0
    answer_results = []
    answer_records = []
    for answer, question in validated_answers:
        is_correct = answer.selected_index == question.correct_index
        if is_correct:
            score += 1
        answer_results.append({
            "question_id": question.id,
            "question_text": question.text,
            "options": _parse_options(question),
            "selected_index": answer.selected_index,
            "correct_index": question.correct_index,
            "is_correct": is_correct,
        })
        answer_records.append(QuizAttemptAnswer(
            question_id=question.id,
            selected_index=answer.selected_index,
            is_correct=is_correct,
        ))
    return score, len(questions), answer_results, answer_records


def _attempt_out(attempt: QuizAttempt, answers=None) -> QuizAttemptOut:
    return QuizAttemptOut(
        id=attempt.id,
        quiz_id=attempt.quiz_id,
        score=attempt.score,
        total=attempt.total,
        submitted_at=attempt.submitted_at,
        answers=[a if isinstance(a, QuizAnswerResultOut) else QuizAnswerResultOut(**a) for a in (answers or [])],
    )


def _teacher_attempt_out(attempt: QuizAttempt, student: Optional[User]) -> QuizTeacherAttemptOut:
    percentage = round(attempt.score / attempt.total * 100, 1) if attempt.total > 0 else 0
    student_name = f"{student.first_name} {student.last_name}" if student else "Unknown"
    return QuizTeacherAttemptOut(
        id=attempt.id,
        quiz_id=attempt.quiz_id,
        student_id=attempt.student_id,
        student_name=student_name,
        score=attempt.score,
        total=attempt.total,
        percentage=percentage,
        submitted_at=attempt.submitted_at,
    )


def _validate_questions(questions, quiz_title: str):
    """Validate question payload and return QuizQuestion objects (uncommitted)."""
    validated = []
    for i, q in enumerate(questions):
        text = q.get('text', '').strip()
        options = [o.strip() for o in q.get('options', []) if o and o.strip()]
        if not text and len(options) < 2:
            continue
        if len(options) < 2:
            raise HTTPException(
                status_code=400,
                detail=f"Quiz '{quiz_title}' question {i + 1} needs at least 2 options",
            )
        correct = q.get('correct', 0)
        if correct < 0 or correct >= len(options):
            raise HTTPException(
                status_code=400,
                detail=f"Quiz '{quiz_title}' question {i + 1} has an invalid correct answer",
            )
        validated.append(QuizQuestion(
            text=text,
            options=json.dumps(options),
            correct_index=correct,
            order_index=len(validated),
        ))
    return validated


@limiter.limit("300/minute")
@router.get("/quizzes", response_model=List[QuizOut])
def list_quizzes(request: Request,
    course_id: str = None,
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """List quizzes scoped to the current user (teacher = their courses)."""
    role = current_user.role.name
    query = db.query(Quiz)

    if role == "teacher":
        from app.routers.courses import get_teacher_course_ids
        teacher_course_ids = get_teacher_course_ids(db, current_user.id, include_inactive=bool(course_id))
        if not teacher_course_ids:
            return []
        query = query.filter(Quiz.course_id.in_(teacher_course_ids))
    elif role == "student":
        student_course_ids = [c.id for c in get_student_courses(db, current_user)]
        if not student_course_ids:
            return []
        query = query.filter(Quiz.course_id.in_(student_course_ids))
        query = query.filter(Quiz.is_published == True)  # noqa: E712
    elif role != "admin":
        raise HTTPException(status_code=403, detail="Access denied")

    if course_id:
        query = query.filter(Quiz.course_id == course_id)

    quizzes = query.order_by(Quiz.created_at.desc()).offset(skip).limit(limit).all()
    return [_quiz_out(q) for q in quizzes]


@limiter.limit("10/minute")
@router.post("/quizzes/parse-file", response_model=ParseFileOut, status_code=status.HTTP_200_OK)
async def parse_quiz_file(request: Request,
    file: UploadFile = File(...),
    current_user: User = Depends(require_teacher),
):
    """Preview-parse a question-bank file (docx/csv/xlsx) without creating a quiz."""
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File too large (max 10MB)")
    questions = parse_question_file(file.filename, content)
    return ParseFileOut(
        questions=[{**q, "correct": q["correct"]} for q in questions],
        count=len(questions),
    )


@limiter.limit("300/minute")
@router.post("/quizzes", response_model=QuizOut, status_code=status.HTTP_201_CREATED)
async def create_quiz(request: Request,
    course_id: str = Form(...),
    title: str = Form(...),
    description: Optional[str] = Form(None),
    time_limit: Optional[int] = Form(None),
    deadline: Optional[str] = Form(None),
    total_questions: Optional[int] = Form(None),
    max_marks: Optional[int] = Form(None),
    questions: Optional[str] = Form(None),
    question_file: Optional[UploadFile] = File(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_teacher),
):
    """Create a quiz from manual questions, an uploaded question bank, or both (teacher only)."""
    course = db.query(Course).filter(Course.id == course_id).first()
    if not course:
        raise HTTPException(status_code=404, detail="Course not found")
    if course.teacher_id != current_user.id:
        raise HTTPException(status_code=403, detail="You can only create quizzes for your courses")
    if not course.is_active:
        raise HTTPException(status_code=400, detail="Cannot create quizzes for an archived course")

    if not title.strip():
        raise HTTPException(status_code=400, detail="Quiz title is required")

    parsed_questions = json.loads(questions) if questions else []

    # Parse uploaded question-bank file (docx/csv/xlsx)
    if question_file and question_file.filename:
        content = await question_file.read()
        if len(content) > 10 * 1024 * 1024:
            raise HTTPException(status_code=400, detail="File too large (max 10MB)")
        file_questions = parse_question_file(question_file.filename, content)
        parsed_questions = list(parsed_questions) + file_questions

    if not parsed_questions:
        raise HTTPException(
            status_code=400,
            detail="A quiz needs at least one question — add questions manually or upload a question-bank file",
        )

    if total_questions is not None and total_questions <= 0:
        raise HTTPException(status_code=400, detail="Total questions must be greater than zero")
    if max_marks is not None and max_marks < 1:
        raise HTTPException(status_code=400, detail="Max marks must be greater than zero")

    parsed_deadline = None
    if deadline:
        parsed_deadline = datetime.fromisoformat(deadline.replace("Z", "+00:00")).replace(tzinfo=None)

    quiz = Quiz(
        course_id=course_id,
        teacher_id=current_user.id,
        title=title.strip(),
        description=description,
        time_limit=time_limit,
        deadline=parsed_deadline,
        total_questions=total_questions,
        max_marks=max_marks,
    )
    db.add(quiz)
    db.flush()

    validated = _validate_questions(parsed_questions, quiz.title)
    if not validated:
        raise HTTPException(status_code=400, detail="No valid questions found")
    if total_questions is not None and total_questions > len(validated):
        raise HTTPException(
            status_code=400,
            detail=f"Total questions ({total_questions}) exceeds the question bank size ({len(validated)})",
        )
    for q in validated:
        q.quiz_id = quiz.id
        db.add(q)

    db.commit()
    db.refresh(quiz)
    return _quiz_out(quiz)


@limiter.limit("300/minute")
@router.post("/quizzes/{quiz_id}/submit", response_model=QuizAttemptOut)
def submit_quiz(request: Request,
    quiz_id: str,
    data: QuizSubmitIn,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Submit quiz answers and return the graded result."""
    if current_user.role.name != "student":
        raise HTTPException(status_code=403, detail="Only students can submit quizzes")

    quiz = db.query(Quiz).filter(Quiz.id == quiz_id).first()
    if not quiz:
        raise HTTPException(status_code=404, detail="Quiz not found")
    if not quiz.is_published:
        raise HTTPException(status_code=400, detail="Quiz is not published yet")
    if quiz.deadline and datetime.now(timezone.utc).replace(tzinfo=None) > quiz.deadline:
        raise HTTPException(status_code=400, detail="Quiz deadline has passed")

    # Check access via session+semester
    course = db.query(Course).filter(Course.id == quiz.course_id).first()
    if not course or not student_has_access(db, current_user, course):
        raise HTTPException(status_code=403, detail="Not enrolled in this course")

    # Check if already attempted
    existing = db.query(QuizAttempt).filter(
        QuizAttempt.quiz_id == quiz_id,
        QuizAttempt.student_id == current_user.id,
    ).first()
    if existing:
        raise HTTPException(status_code=400, detail="You have already attempted this quiz")

    # Enforce the countdown server-side (grace absorbs network latency)
    if quiz.time_limit:
        timer = db.query(QuizTimer).filter(
            QuizTimer.quiz_id == quiz_id,
            QuizTimer.student_id == current_user.id,
        ).first()
        if timer:
            expires = timer.started_at + timedelta(
                minutes=quiz.time_limit, seconds=TIMER_GRACE_SECONDS
            )
            if datetime.now(timezone.utc).replace(tzinfo=None) > expires:
                raise HTTPException(status_code=400, detail="Time is up — the quiz timer has expired")

    # Grade against the questions assigned to this student (same deterministic
    # sample/shuffle the student saw when fetching the quiz)
    assigned = _assigned_questions(quiz, current_user.id)
    score, total, answer_results, answer_records = _score_answers(
        assigned, data.answers, require_all=not data.timed_out
    )

    attempt = QuizAttempt(
        quiz_id=quiz_id,
        student_id=current_user.id,
        score=score,
        total=total,
        grading_status="graded",
    )
    db.add(attempt)
    try:
        db.flush()
        for answer in answer_records:
            answer.attempt_id = attempt.id
            db.add(answer)
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=400, detail="You have already attempted this quiz")
    db.refresh(attempt)

    return _attempt_out(attempt, answer_results)


@limiter.limit("300/minute")
@router.get("/quizzes/{quiz_id}/attempts", response_model=Optional[QuizAttemptOut])
def get_my_attempt(request: Request,
    quiz_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Get the current student's attempt result for a quiz."""
    attempt = db.query(QuizAttempt).filter(
        QuizAttempt.quiz_id == quiz_id,
        QuizAttempt.student_id == current_user.id,
    ).first()
    if not attempt:
        return None

    answers_out = []
    question_ids = list({aa.question_id for aa in attempt.answers})
    questions = {q.id: q for q in db.query(QuizQuestion).filter(QuizQuestion.id.in_(question_ids)).all()} if question_ids else {}
    for aa in attempt.answers:
        q = questions.get(aa.question_id)
        if q:
            answers_out.append(QuizAnswerResultOut(
                question_id=q.id,
                question_text=q.text,
                options=_parse_options(q),
                selected_index=aa.selected_index,
                correct_index=q.correct_index,
                is_correct=aa.is_correct,
            ))

    return _attempt_out(attempt, answers_out)


@limiter.limit("300/minute")
@router.get("/quizzes/{quiz_id}/all-attempts", response_model=List[QuizTeacherAttemptOut])
def get_all_attempts(request: Request,
    quiz_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_teacher),
):
    """Get all student attempts for a quiz (teacher only)."""
    quiz = db.query(Quiz).filter(Quiz.id == quiz_id).first()
    if not quiz:
        raise HTTPException(status_code=404, detail="Quiz not found")

    if quiz.course and quiz.course.teacher_id != current_user.id:
        raise HTTPException(status_code=403, detail="Access denied")

    attempts = db.query(QuizAttempt).filter(QuizAttempt.quiz_id == quiz_id).all()
    student_ids = list({a.student_id for a in attempts})
    students = {u.id: u for u in db.query(User).filter(User.id.in_(student_ids)).all()} if student_ids else {}
    return [_teacher_attempt_out(a, students.get(a.student_id)) for a in attempts]


def _export_filename(quiz: Quiz) -> str:
    course = quiz.course
    raw = f"{(course.course_code if course else '') or 'quiz'}_{quiz.title or 'results'}_results.xlsx"
    return re.sub(r"[^\w\-. ]+", "_", raw)[:100]


@limiter.limit("300/minute")
@router.get("/quizzes/{quiz_id}/export")
def export_quiz_results(request: Request,
    quiz_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_teacher),
):
    """Download all results for a quiz as an Excel file (owner teacher only)."""
    quiz = db.query(Quiz).filter(Quiz.id == quiz_id).first()
    if not quiz:
        raise HTTPException(status_code=404, detail="Quiz not found")
    if quiz.course and quiz.course.teacher_id != current_user.id:
        raise HTTPException(status_code=403, detail="Access denied")

    attempts = (db.query(QuizAttempt)
                 .filter(QuizAttempt.quiz_id == quiz_id)
                 .order_by(QuizAttempt.submitted_at)
                 .all())
    student_ids = list({a.student_id for a in attempts})
    students = {u.id: u for u in db.query(User).filter(User.id.in_(student_ids)).all()} if student_ids else {}
    profiles = ({p.user_id: p for p in db.query(StudentProfile).filter(StudentProfile.user_id.in_(student_ids)).all()}
                if student_ids else {})

    course = quiz.course
    amber_fill = PatternFill("solid", fgColor="F59E0B")
    navy_fill = PatternFill("solid", fgColor="1F2A44")
    light_fill = PatternFill("solid", fgColor="E9EEF6")
    band_fill = PatternFill("solid", fgColor="F7F9FC")
    thin = Side(style="thin", color="D3DCE8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    for i, w in enumerate([6, 26, 15, 30, 9, 9, 13, 21], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    def styled_merged_row(row, value, fill, font, height=None):
        for c in range(1, 9):
            ws.cell(row=row, column=c).fill = fill
        cell = ws.cell(row=row, column=1, value=value)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell.font = font
        cell.alignment = left_align
        if height:
            ws.row_dimensions[row].height = height

    # Amber accent strip + navy title + light info rows
    styled_merged_row(1, None, amber_fill, Font(size=6), height=6)
    styled_merged_row(2, f"Quiz Results — {quiz.title}", navy_fill,
                      Font(bold=True, size=14, color="FFFFFF"), height=26)
    styled_merged_row(3, f"Course: {f'{course.course_code} — {course.title}' if course else ''}",
                      light_fill, Font(bold=True, size=11, color="1F2A44"), height=18)
    styled_merged_row(4, f"Time limit: {quiz.time_limit or 'No limit'} min | Deadline: {quiz.deadline or 'None'} | Attempts: {len(attempts)}",
                      light_fill, Font(italic=True, size=10, color="475569"), height=16)
    ws.row_dimensions[5].height = 6

    # Column headers
    headers = ["#", "Student", "Roll Number", "Email", "Score", "Total", "Percentage", "Submitted At"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=6, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = navy_fill
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[6].height = 20

    # Attempt rows (banded, bordered, colored percentages)
    r = 7
    for i, a in enumerate(attempts, 1):
        u = students.get(a.student_id)
        p = profiles.get(a.student_id)
        pct = round(a.score / a.total * 100, 1) if a.total else 0
        submitted = a.submitted_at
        values = [
            i,
            f"{u.first_name} {u.last_name}".strip() if u else "Unknown",
            (p.roll_number if p and p.roll_number else ""),
            (u.email if u else ""),
            a.score,
            a.total,
            pct,
            submitted if isinstance(submitted, datetime) else str(submitted or ""),
        ]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = border
            cell.alignment = left_align if c in (2, 4) else center
            if i % 2 == 0:
                cell.fill = band_fill
        if isinstance(submitted, datetime):
            ws.cell(row=r, column=8).number_format = "yyyy-mm-dd hh:mm"
        pct_color = "059669" if pct >= 70 else "D97706" if pct >= 40 else "DC2626"
        ws.cell(row=r, column=7).font = Font(bold=True, color=pct_color)
        r += 1

    ws.freeze_panes = "A7"

    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{_export_filename(quiz)}"'},
    )


@limiter.limit("300/minute")
@router.get("/quizzes/{quiz_id}", response_model=QuizDetailOut)
def get_quiz(request: Request,
    quiz_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Get a quiz with its questions.

    Students receive their own shuffled subset (no correct answers);
    the owning teacher receives the full bank with correct answers.
    """
    quiz = db.query(Quiz).filter(Quiz.id == quiz_id).first()
    if not quiz:
        raise HTTPException(status_code=404, detail="Quiz not found")

    role = current_user.role.name
    include_correct = False
    started_at = None
    if role == "teacher":
        if not quiz.course or quiz.course.teacher_id != current_user.id:
            raise HTTPException(status_code=403, detail="Access denied")
        include_correct = True
        assigned = list(quiz.questions)
    elif role == "student":
        if not quiz.is_published:
            raise HTTPException(status_code=404, detail="Quiz not found")
        course = db.query(Course).filter(Course.id == quiz.course_id).first()
        if not course or not student_has_access(db, current_user, course):
            raise HTTPException(status_code=403, detail="Access denied")
        assigned = _assigned_questions(quiz, current_user.id)
        if quiz.time_limit:
            started_at = _ensure_quiz_timer(db, quiz, current_user.id)
    elif role == "admin":
        assigned = list(quiz.questions)
        include_correct = True
    else:
        raise HTTPException(status_code=403, detail="Access denied")

    out = _quiz_out(quiz)
    return QuizDetailOut(
        **out.model_dump(),
        questions=[_question_out(q, include_correct, order=i) for i, q in enumerate(assigned)],
        started_at=started_at,
    )


@limiter.limit("300/minute")
@router.delete("/quizzes/{quiz_id}")
def delete_quiz(request: Request,
    quiz_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_teacher),
):
    """Delete a quiz and its questions (teacher only, must own)."""
    quiz = db.query(Quiz).filter(Quiz.id == quiz_id).first()
    if not quiz:
        raise HTTPException(status_code=404, detail="Quiz not found")
    if not quiz.course or quiz.course.teacher_id != current_user.id:
        raise HTTPException(status_code=403, detail="Access denied")

    db.delete(quiz)
    db.commit()
    return {"message": "Quiz deleted successfully"}
